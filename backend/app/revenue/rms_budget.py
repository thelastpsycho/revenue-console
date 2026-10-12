"""Operator-entered, segment-level daily revenue/room-night budget.

Source: an annual Market Segment Budget workbook (finance-produced). Re-run
the importer once per year when a new workbook arrives - it's idempotent per
(month, market_id, basis). Imports the workbook's "SUM BUDGET" sheet (monthly
room nights/ARR/revenue per market segment), then pro-rates each month evenly
across its days into `daily_budget`, keyed by the same `market_id` pace and
production already filter on, so a budget comparison needs no new taxonomy.

The workbook's segment labels don't match the PMS `market` field's spelling,
but both taxonomies have the same 10 segments - confirmed against
booking_events.db's distinct (market_id, market) pairs.
"""
import sqlite3
from calendar import monthrange
from datetime import date

import openpyxl

from ..infrastructure.paths import get_data_path

DB_FILE = "rms_budget.db"
SHEET_NAME = "SUM BUDGET"
BASES = ("gross", "net")

MONTH_NUMBERS = {
    "january": 1, "february": 2, "march": 3, "april": 4, "may": 5, "june": 6,
    "july": 7, "august": 8, "september": 9, "october": 10, "november": 11, "december": 12,
}

MARKET_ID_BY_SEGMENT = {
    "fit - direct": 3,               # PMS market "FIT"
    "fit - travel agent": 4,         # PMS market "Travel Agent"
    "corporate": 2,                  # PMS market "Corporate"
    "govt": 5,                       # PMS market "Government"
    "group/it series": 6,            # PMS market "Group - IT Series"
    "adhoc": 11,                     # PMS market "Group - AdHoc"
    "meetings": 7,                   # PMS market "Meeting"
    "package": 9,                    # PMS market "Package"
    "online travel agent (ota)": 8,  # PMS market "Online"
    "unique": 10,                    # PMS market "Unique"
}


def _connect():
    conn = sqlite3.connect(get_data_path(DB_FILE), timeout=30)
    conn.row_factory = sqlite3.Row
    conn.executescript("""
        CREATE TABLE IF NOT EXISTS market_segment_budgets (
            month TEXT NOT NULL,
            market_id INTEGER NOT NULL,
            segment TEXT NOT NULL,
            basis TEXT NOT NULL,
            room_nights REAL NOT NULL,
            arr REAL NOT NULL,
            revenue REAL NOT NULL,
            PRIMARY KEY (month, market_id, basis)
        );
        CREATE TABLE IF NOT EXISTS daily_budget (
            date TEXT NOT NULL,
            market_id INTEGER NOT NULL,
            basis TEXT NOT NULL,
            rooms REAL NOT NULL,
            revenue REAL NOT NULL,
            PRIMARY KEY (date, market_id, basis)
        );
        CREATE INDEX IF NOT EXISTS idx_daily_budget_date ON daily_budget(date);
    """)
    return conn


def _segment_key(label):
    return " ".join((label or "").split()).lower()


def read_sum_budget_sheet(path, year):
    """[(month YYYY-MM, market_id, segment, room_nights, arr, revenue), ...]
    from the workbook's "SUM BUDGET" sheet. Raises on an unmapped segment
    label so a future workbook with a renamed/added segment fails loudly
    instead of silently dropping budget."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[SHEET_NAME]
        out = []
        for r in range(2, ws.max_row + 1):
            month_name, segment, _share, rn, arr, rev = (ws.cell(row=r, column=c).value for c in range(1, 7))
            if not month_name or not segment:
                continue  # the sheet's trailing "Total" row has no segment
            month_num = MONTH_NUMBERS.get(month_name.strip().lower())
            if month_num is None:
                continue
            key = _segment_key(segment)
            if key not in MARKET_ID_BY_SEGMENT:
                raise ValueError(f"Unmapped budget segment {segment!r} in {path}")
            out.append((f"{year:04d}-{month_num:02d}", MARKET_ID_BY_SEGMENT[key], segment.strip(),
                       float(rn or 0), float(arr or 0), float(rev or 0)))
        return out
    finally:
        wb.close()


def materialize_daily(conn, months=None):
    """Rebuilds `daily_budget` from `market_segment_budgets` by pro-rating
    each month's segment totals evenly across its days. Pass the just-
    imported months to limit the rebuild; omit to rebuild everything."""
    query = "SELECT * FROM market_segment_budgets"
    params = []
    if months:
        query += f" WHERE month IN ({','.join('?' * len(months))})"
        params = list(months)
    rows = conn.execute(query, params).fetchall()
    for month in sorted({row["month"] for row in rows}):
        conn.execute("DELETE FROM daily_budget WHERE date BETWEEN ? AND ?", (f"{month}-01", f"{month}-31"))
    for row in rows:
        year, month_num = (int(part) for part in row["month"].split("-"))
        days_in_month = monthrange(year, month_num)[1]
        daily_rooms = row["room_nights"] / days_in_month
        daily_revenue = row["revenue"] / days_in_month
        for day in range(1, days_in_month + 1):
            conn.execute(
                "INSERT OR REPLACE INTO daily_budget (date, market_id, basis, rooms, revenue) VALUES (?, ?, ?, ?, ?)",
                (date(year, month_num, day).isoformat(), row["market_id"], row["basis"], daily_rooms, daily_revenue))
    conn.commit()


def daily_totals(start, end, market_id=None, basis="gross"):
    """{date: {'rooms':, 'revenue':}} within [start, end], summed across
    segments when market_id is empty. {} if nothing has been imported yet."""
    conn = _connect()
    try:
        query = "SELECT date, SUM(rooms) rooms, SUM(revenue) revenue FROM daily_budget WHERE date BETWEEN ? AND ? AND basis = ?"
        params = [start, end, basis]
        if market_id not in (None, ""):
            query += " AND market_id = ?"
            params.append(int(market_id))
        rows = conn.execute(query + " GROUP BY date", params).fetchall()
        return {row["date"]: {"rooms": row["rooms"], "revenue": row["revenue"]} for row in rows}
    finally:
        conn.close()


def segment_totals(start, end, basis="gross"):
    """{market_id (as str): {'rooms':, 'revenue':}} within [start, end], one
    entry per market segment with imported budget data."""
    conn = _connect()
    try:
        rows = conn.execute(
            "SELECT market_id, SUM(rooms) rooms, SUM(revenue) revenue FROM daily_budget "
            "WHERE date BETWEEN ? AND ? AND basis = ? GROUP BY market_id", (start, end, basis)).fetchall()
        return {str(row["market_id"]): {"rooms": row["rooms"], "revenue": row["revenue"]} for row in rows}
    finally:
        conn.close()


def import_budget(path, year, basis="gross"):
    """Imports one workbook into `market_segment_budgets` and rebuilds the
    affected months of `daily_budget`. Safe to re-run: replaces existing rows
    for the same (month, market_id, basis). Returns the number of
    month/segment rows imported."""
    if basis not in BASES:
        raise ValueError(f"basis must be one of {BASES}")
    records = read_sum_budget_sheet(path, year)
    conn = _connect()
    try:
        for month, market_id, segment, rn, arr, rev in records:
            conn.execute("""
                INSERT OR REPLACE INTO market_segment_budgets
                    (month, market_id, segment, basis, room_nights, arr, revenue)
                VALUES (?, ?, ?, ?, ?, ?, ?)""",
                (month, market_id, segment, basis, rn, arr, rev))
        conn.commit()
        materialize_daily(conn, sorted({month for month, *_ in records}))
        return len(records)
    finally:
        conn.close()


if __name__ == "__main__":  # python -m app.revenue.rms_budget <xlsx-path> <year> [gross|net]
    import sys

    xlsx_path, budget_year = sys.argv[1], int(sys.argv[2])
    budget_basis = sys.argv[3] if len(sys.argv) > 3 else "gross"
    imported = import_budget(xlsx_path, budget_year, budget_basis)
    print(f"Imported {imported} segment-month rows for {budget_year} ({budget_basis}) from {xlsx_path}")
